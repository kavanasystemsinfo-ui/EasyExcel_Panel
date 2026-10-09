"""Contrato de exportación de hoja a .xlsx (Fase 5).

TDD: estos tests definen el comportamiento antes de implementar.
"""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.core.config import settings
from app.main import app

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client: TestClient, data: bytes, filename: str = "demo.xlsx"):
    return client.post(
        "/api/v1/workbooks",
        files={"file": (filename, data, CONTENT_TYPE)},
    )


@pytest.fixture(autouse=True)
def uploads_en_tmp(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "upload_dir", tmp_path)


@pytest.fixture()
def client():
    return TestClient(app)


@pytest.fixture()
def xlsx_2_hojas() -> bytes:
    return make_xlsx(
        {
            "Ventas 2026": [
                ["Producto", "Unidades", "Importe"],
                ["Sofa", 3, 450.5],
                ["Mesa", 1, 120.0],
                ["Silla", 10, 300.0],
            ],
            "Resumen": [
                ["Concepto", "Valor"],
                ["Total", 870.5],
            ],
        }
    )


def export(client: TestClient, workbook_id: str, payload: dict):
    return client.post(
        f"/api/v1/workbooks/{workbook_id}/sheets/Ventas 2026/export",
        json=payload,
    )


# ------------------------------------------------------------------ export


def test_export_devuelve_xlsx_con_lo_enviado(client, xlsx_2_hojas):
    wid = upload(client, xlsx_2_hojas).json()["id"]
    resp = export(
        client,
        wid,
        {
            "header": ["Producto", "Importe"],
            "rows": [["Mesa", 120.0], ["Silla", 300.0]],
        },
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(CONTENT_TYPE)
    assert "attachment" in resp.headers["content-disposition"]
    assert "Ventas" in resp.headers["content-disposition"]
    wb = load_workbook(BytesIO(resp.content))
    ws = wb.active
    assert list(ws.values) == [("Producto", "Importe"), ("Mesa", 120.0), ("Silla", 300.0)]


def test_export_con_solo_cabecera(client, xlsx_2_hojas):
    wid = upload(client, xlsx_2_hojas).json()["id"]
    resp = export(client, wid, {"header": ["Producto", "Unidades"], "rows": []})
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.content))
    assert list(wb.active.values) == [("Producto", "Unidades")]


def test_export_admite_nulos_y_texto_largo(client, xlsx_2_hojas):
    wid = upload(client, xlsx_2_hojas).json()["id"]
    resp = export(
        client,
        wid,
        {"header": ["Nota"], "rows": [[None], ['texto con "comillas" y ñ']]},
    )
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.content))
    assert list(wb.active.values) == [
        ("Nota",),
        (None,),
        ('texto con "comillas" y ñ',),
    ]


def test_export_hoja_inexistente_404(client, xlsx_2_hojas):
    wid = upload(client, xlsx_2_hojas).json()["id"]
    resp = client.post(
        f"/api/v1/workbooks/{wid}/sheets/NoExiste/export",
        json={"header": ["A"], "rows": [[1]]},
    )
    assert resp.status_code == 404


def test_export_libro_inexistente_404(client):
    resp = client.post(
        "/api/v1/workbooks/no-existe/sheets/Ventas/export",
        json={"header": ["A"], "rows": [[1]]},
    )
    assert resp.status_code == 404


def test_export_header_vacio_400(client, xlsx_2_hojas):
    wid = upload(client, xlsx_2_hojas).json()["id"]
    resp = export(client, wid, {"header": [], "rows": []})
    assert resp.status_code == 400


def test_export_demasiadas_filas_422(client, xlsx_2_hojas, monkeypatch):
    wid = upload(client, xlsx_2_hojas).json()["id"]
    monkeypatch.setattr(settings, "max_rows_per_sheet", 2)
    resp = export(client, wid, {"header": ["A"], "rows": [[1], [2], [3]]})
    assert resp.status_code == 422


def test_export_demasiadas_columnas_422(client, xlsx_2_hojas, monkeypatch):
    wid = upload(client, xlsx_2_hojas).json()["id"]
    monkeypatch.setattr(settings, "max_columns", 2)
    resp = export(client, wid, {"header": ["A", "B", "C"], "rows": []})
    assert resp.status_code == 422


def test_export_fila_mas_ancha_que_header_422(client, xlsx_2_hojas):
    wid = upload(client, xlsx_2_hojas).json()["id"]
    resp = export(client, wid, {"header": ["A"], "rows": [[1, 2]]})
    assert resp.status_code == 422


def test_export_demasiadas_celdas_422(client, xlsx_2_hojas, monkeypatch):
    wid = upload(client, xlsx_2_hojas).json()["id"]
    monkeypatch.setattr(settings, "max_export_cells", 4)
    resp = export(client, wid, {"header": ["A", "B"], "rows": [[1, 2], [3, 4], [5, 6]]})
    assert resp.status_code == 422

