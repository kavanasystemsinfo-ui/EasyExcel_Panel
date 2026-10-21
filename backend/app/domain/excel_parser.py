"""Lectura de Excel con timeout (ADR-0001). Compatible con serverless (sin multiprocessing).

Formatos de entrada: .xlsx (openpyxl), .xls (xlrd), .ods (odfpy), .csv (stdlib).
La exportación siempre genera .xlsx.
"""

import csv as csv_mod
import re
import signal
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from odf import table as odf_table
from odf import teletype
from odf.opendocument import load as ods_load
from openpyxl import Workbook, load_workbook
import xlrd

from app.core.config import settings
from app.errors import (
    InvalidWorkbookError,
    WorkbookLimitError,
    WorkbookNotFoundError,
    WorkbookParseTimeoutError,
)

FORMATOS_IMPORT = ("xlsx", "xls", "ods", "csv")


class TimeoutError(Exception):
    """Timeout personalizado para el parser."""


def _timeout_handler(signum, frame):
    raise TimeoutError("Timeout")


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float, str, bool)):
        return value
    return str(value)


def _width(row: tuple[Any, ...]) -> int:
    width = len(row)
    while width > 0 and row[width - 1] is None:
        width -= 1
    return width


# --------------------------------------------------- formatos no xlsx
#
# El storage guarda siempre book.xlsx, asi que el formato se detecta por
# magic bytes y el nombre de la hoja del csv viene del nombre original.


OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _detect_format(path: str) -> str:
    import zipfile

    with open(path, "rb") as fh:
        head = fh.read(8)
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(path) as archive:
                if "mimetype" in archive.namelist():
                    if b"opendocument.spreadsheet" in archive.read("mimetype"):
                        return "ods"
        except zipfile.BadZipFile:
            pass
        return "xlsx"
    if head.startswith(OLE2_MAGIC):
        return "xls"
    return "csv"


def _csv_number(valor: str, coma_decimal: bool) -> Any:
    texto = valor.strip()
    if not texto:
        return ""
    if re.fullmatch(r"-?\d+", texto):
        return int(texto)
    if coma_decimal and re.fullmatch(r"-?\d+,\d+", texto):
        return float(texto.replace(",", "."))
    if re.fullmatch(r"-?\d+\.\d+", texto):
        return float(texto)
    return valor


def _load_csv(path: str, nombre: str | None) -> list[dict]:
    data = Path(path).read_bytes()
    texto: str | None = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if texto is None:
        texto = data.decode("latin-1", errors="replace")
    primera = next((linea for linea in texto.splitlines() if linea.strip()), "")
    delimitador = max((";", ",", "\t"), key=primera.count)
    coma_decimal = delimitador != ","
    max_filas = settings.max_rows_per_sheet + 2
    max_cols = settings.max_columns + 1
    filas: list[list[Any]] = []
    for fila in csv_mod.reader(StringIO(texto), delimiter=delimitador):
        filas.append([_csv_number(c, coma_decimal) for c in fila[:max_cols]])
        if len(filas) >= max_filas:
            break
    hoja = Path(nombre).stem if nombre else "libro"
    return [{"name": hoja or "libro", "rows": filas}]


def _load_xls(path: str) -> list[dict]:
    book = xlrd.open_workbook(path)
    max_filas = settings.max_rows_per_sheet + 2
    max_cols = settings.max_columns + 1
    sheets: list[dict] = []
    for sh in book.sheets():
        filas: list[list[Any]] = []
        for r in range(min(sh.nrows, max_filas)):
            fila: list[Any] = []
            for c in range(min(sh.ncols, max_cols)):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    dt = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                    fila.append(dt.isoformat())
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    value = cell.value
                    fila.append(int(value) if value == int(value) else value)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    fila.append(bool(cell.value))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    fila.append("")
                else:
                    fila.append(str(cell.value))
            filas.append(fila)
        sheets.append({"name": sh.name, "rows": filas})
    return sheets


def _ods_cell_value(node) -> Any:
    # odfpy expone las "clases" como factories: distinguir por qname.
    if node.qname[1] == "covered-table-cell":
        return ""
    valuetype = node.getAttribute("valuetype")
    if valuetype in ("float", "percentage", "currency"):
        raw = node.getAttribute("value")
        try:
            number = float(raw)
            return int(number) if number == int(number) else number
        except (TypeError, ValueError):
            return teletype.extractText(node)
    if valuetype == "boolean":
        return node.getAttribute("booleanvalue") == "true"
    if valuetype == "date":
        return node.getAttribute("datevalue") or teletype.extractText(node)
    return teletype.extractText(node)


def _load_ods(path: str) -> list[dict]:
    doc = ods_load(path)
    max_filas = settings.max_rows_per_sheet + 2
    max_cols = settings.max_columns + 1
    sheets: list[dict] = []
    for tbl in doc.spreadsheet.getElementsByType(odf_table.Table):
        name = tbl.getAttribute("name") or "Hoja"
        filas: list[list[Any]] = []
        for row in tbl.getElementsByType(odf_table.TableRow):
            repeticiones = int(row.getAttribute("numberrowsrepeated") or 1)
            celdas: list[Any] = []
            for node in row.childNodes:
                qname = getattr(node, "qname", None)
                if qname is None or qname[1] not in ("table-cell", "covered-table-cell"):
                    continue
                repeat = int(node.getAttribute("numbercolumnsrepeated") or 1)
                value = _ods_cell_value(node)
                for _ in range(min(repeat, max_cols - len(celdas))):
                    celdas.append(value)
            # LibreOffice rellena el final con filas vacias repetidas:
            # no materializar ese padding (dispararia falsos limites).
            if repeticiones > 1 and all(c == "" for c in celdas):
                continue
            for _ in range(min(repeticiones, max_filas - len(filas))):
                filas.append(celdas)
            if len(filas) >= max_filas:
                break
        sheets.append({"name": name, "rows": filas})
    return sheets


def _load_sheets(path: str, fmt: str, nombre: str | None) -> list[dict]:
    if fmt == "csv":
        return _load_csv(path, nombre)
    if fmt == "xls":
        return _load_xls(path)
    if fmt == "ods":
        return _load_ods(path)
    raise InvalidWorkbookError(f"Formato no soportado: .{fmt}")


def _metadata_from_rows(sheets: list[dict]) -> list[dict]:
    out: list[dict] = []
    for sheet in sheets:
        rows = sheet["rows"]
        cols = max((_width(tuple(row)) for row in rows), default=0)
        out.append({"name": sheet["name"], "rows": max(len(rows) - 1, 0), "cols": cols})
    return out


def _read_rows_from_sheets(
    sheets: list[dict], sheet: str, offset: int, limit: int
) -> dict:
    match = next((s for s in sheets if s["name"] == sheet), None)
    if match is None:
        raise WorkbookNotFoundError(f"La hoja '{sheet}' no existe")
    rows = match["rows"]
    header = [_cell(value) for value in rows[0]] if rows else []
    data = rows[1:]
    page = [[_cell(value) for value in row] for row in data[offset : offset + limit]]
    return {"header": header, "rows": page, "total": len(data)}


def _extract_metadata_sync(path: str) -> list[dict]:
    """Extrae metadatos sin multiprocessing."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict] = []
    for ws in workbook.worksheets:
        total = 0
        cols = 0
        for row in ws.iter_rows(values_only=True):
            total += 1
            cols = max(cols, _width(row))
        sheets.append({"name": ws.title, "rows": max(total - 1, 0), "cols": cols})
    workbook.close()
    return sheets


def _read_rows_sync(path: str, sheet: str, offset: int, limit: int) -> dict:
    """Lee filas sin multiprocessing."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet not in workbook.sheetnames:
        workbook.close()
        raise WorkbookNotFoundError(f"La hoja '{sheet}' no existe")
    ws = workbook[sheet]
    header: list[Any] = []
    page: list[list[Any]] = []
    total = 0
    seen_header = False
    for row in ws.iter_rows(values_only=True):
        if not seen_header:
            header = [_cell(value) for value in row]
            seen_header = True
            continue
        if total >= offset and len(page) < limit:
            page.append([_cell(value) for value in row])
        total += 1
    workbook.close()
    return {"header": header, "rows": page, "total": total}


def _run_with_timeout(func, *args, timeout_s: float):
    """Ejecuta una función con timeout usando signal (solo Unix, funciona en Vercel)."""
    # En Windows/serverless sin signal, ejecutar sin timeout
    try:
        old_handler = signal.signal(signal.SIGALRM, _timeout_handler)
        signal.setitimer(signal.ITIMER_REAL, timeout_s)
        try:
            result = func(*args)
        finally:
            signal.setitimer(signal.ITIMER_REAL, 0)
            signal.signal(signal.SIGALRM, old_handler)
        return result
    except (AttributeError, ValueError, OSError):
        # signal no disponible (Windows, o entorno restringido) -> sin timeout
        return func(*args)
    except TimeoutError:
        raise WorkbookParseTimeoutError(f"El parser supero {timeout_s} segundos")


def _check_limits(sheets: list[dict]) -> None:
    if len(sheets) > settings.max_sheets:
        raise WorkbookLimitError(f"Maximo {settings.max_sheets} hojas por libro")
    for sheet in sheets:
        if sheet["cols"] > settings.max_columns:
            raise WorkbookLimitError(
                f"La hoja '{sheet['name']}' supera {settings.max_columns} columnas"
            )
        if sheet["rows"] > settings.max_rows_per_sheet:
            raise WorkbookLimitError(
                f"La hoja '{sheet['name']}' supera {settings.max_rows_per_sheet} filas de datos"
            )


def _extract_metadata_generic(path: str, fmt: str, nombre: str | None) -> list[dict]:
    return _metadata_from_rows(_load_sheets(path, fmt, nombre))


def _read_rows_generic(
    path: str, fmt: str, nombre: str | None, sheet: str, offset: int, limit: int
) -> dict:
    return _read_rows_from_sheets(_load_sheets(path, fmt, nombre), sheet, offset, limit)


def extract_metadata(path: Path, nombre: str | None = None) -> list[dict]:
    fmt = _detect_format(str(path))
    if fmt == "xlsx":
        sheets = _run_with_timeout(
            _extract_metadata_sync, str(path), timeout_s=settings.parse_timeout_s
        )
    else:
        sheets = _run_with_timeout(
            _extract_metadata_generic,
            str(path), fmt, nombre,
            timeout_s=settings.parse_timeout_s,
        )
    _check_limits(sheets)
    return sheets


def read_rows(
    path: Path, sheet: str, offset: int, limit: int, nombre: str | None = None
) -> dict:
    fmt = _detect_format(str(path))
    if fmt == "xlsx":
        return _run_with_timeout(
            _read_rows_sync,
            str(path), sheet, offset, limit,
            timeout_s=settings.parse_timeout_s,
        )
    return _run_with_timeout(
        _read_rows_generic,
        str(path), fmt, nombre, sheet, offset, limit,
        timeout_s=settings.parse_timeout_s,
    )


# ------------------------------------------------------------------ export

_EXPORT_CELL_LIMIT = 32_767


def _export_scalar(value: Any) -> Any:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = value if isinstance(value, str) else str(value)
    return text[:_EXPORT_CELL_LIMIT]


def build_export(header: list[str], rows: list[list[Any]]) -> bytes:
    if not header:
        raise InvalidWorkbookError("La exportacion necesita al menos una columna")
    if len(header) > settings.max_columns:
        raise WorkbookLimitError(f"Maximo {settings.max_columns} columnas por hoja")
    if len(rows) > settings.max_rows_per_sheet:
        raise WorkbookLimitError(f"Maximo {settings.max_rows_per_sheet} filas por hoja")
    if len(rows) * len(header) > settings.max_export_cells:
        raise WorkbookLimitError(
            f"Maximo {settings.max_export_cells} celdas por exportacion"
        )
    for index, row in enumerate(rows, start=2):
        if len(row) > len(header):
            raise WorkbookLimitError(
                f"La fila {index} tiene {len(row)} valores "
                f"y la cabecera {len(header)} columnas"
            )
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([_export_scalar(cell) for cell in header])
    for row in rows:
        sheet.append([_export_scalar(cell) for cell in row])
    buf = BytesIO()
    workbook.save(buf)
    return buf.getvalue()
