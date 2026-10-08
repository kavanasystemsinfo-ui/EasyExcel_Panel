import logging

from fastapi import FastAPI, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse

from backend.app.ai import assistant
from backend.app.api.routes import copiloto, health, workbooks
from backend.app.core.config import settings
from backend.app.errors import (
    InvalidWorkbookError,
    WorkbookLimitError,
    WorkbookNotFoundError,
    WorkbookParseTimeoutError,
    WorkbookTooLargeError,
)


def create_app() -> FastAPI:
    app = FastAPI(title=settings.app_name, version=settings.version)
    if not assistant.logger.handlers:
        log_handler = logging.StreamHandler()
        log_handler.setFormatter(logging.Formatter("%(levelname)s %(name)s: %(message)s"))
        assistant.logger.addHandler(log_handler)
        assistant.logger.setLevel(logging.INFO)
    app.add_middleware(
        CORSMiddleware,
        allow_origins=settings.cors_origins,
        allow_credentials=True,
        allow_methods=["GET", "POST", "DELETE"],
        allow_headers=["*"],
    )
    app.include_router(health.router)
    app.include_router(workbooks.router)
    app.include_router(copiloto.router)

    def handler(status_code: int):
        def _handle(request: Request, exc: Exception) -> JSONResponse:
            detail = str(exc) or type(exc).__name__
            return JSONResponse({"detail": detail}, status_code=status_code)

        return _handle

    app.add_exception_handler(InvalidWorkbookError, handler(400))
    app.add_exception_handler(WorkbookTooLargeError, handler(413))
    app.add_exception_handler(WorkbookNotFoundError, handler(404))
    app.add_exception_handler(WorkbookLimitError, handler(422))
    app.add_exception_handler(WorkbookParseTimeoutError, handler(422))
    return app


app = create_app()

@app.get("/api/v1/demo-data")
async def get_demo_data():
    """Endpoint que devuelve los datos del Excel por defecto"""
    import os
    from openpyxl import load_workbook
    from fastapi import HTTPException

    default_path = os.path.join(os.path.dirname(__file__), "..", "data", "default_demo.xlsx")
    if not os.path.exists(default_path):
        # Gather debug info
        data_dir = os.path.join(os.path.dirname(__file__), "..", "data")
        parent_dir = os.path.dirname(__file__)
        debug = {
            "path": default_path,
            "exists": False,
            "parent_dir": parent_dir,
            "data_dir": data_dir,
            "data_dir_contents": os.listdir(data_dir) if os.path.exists(data_dir) else "data_dir not found",
            "parent_dir_contents": os.listdir(parent_dir) if os.path.exists(parent_dir) else "parent_dir not found"
        }
        raise HTTPException(status_code=404, detail=debug)

    wb = load_workbook(default_path, data_only=True)
    sheet = wb.active
    data = []
    for row in sheet.iter_rows(values_only=True):
        data.append(list(row))
    if not data:
        raise HTTPException(status_code=404, detail="Demo file is empty")
    headers = data[0]
    rows = data[1:]
    return {"headers": headers, "rows": rows}
