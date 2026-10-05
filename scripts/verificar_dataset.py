#!/usr/bin/env python3
"""Read-back independiente de data/easyexcel_demo.xlsx: relee el archivo
escrito y revalida las reglas criticas sin pasar por el generador."""

import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

HOY = date(2026, 10, 5)
SMI = 1221.0
BASE = Path(__file__).resolve().parent.parent
SRC = BASE / "data" / "easyexcel_demo.xlsx"

wb = load_workbook(SRC, data_only=True)
esperadas = ["Info", "Empleados", "Centros", "Asignaciones", "Contratos",
             "Vacaciones", "Proveedores", "Auditorias", "Maquinaria", "Revisiones"]
errores = []

if wb.sheetnames != esperadas:
    errores.append(f"hojas: {wb.sheetnames} != {esperadas}")


def filas(nombre):
    ws = wb[nombre]
    datos = list(ws.iter_rows(values_only=True))
    cab = datos[0]
    return [dict(zip(cab, f)) for f in datos[1:]], cab


def como_fecha(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date(*reversed([int(x) for x in str(v).split("/")]))


empleados, _ = filas("Empleados")
centros, _ = filas("Centros")
asignaciones, _ = filas("Asignaciones")
contratos, _ = filas("Contratos")
vacaciones, _ = filas("Vacaciones")
proveedores, _ = filas("Proveedores")
auditorias, _ = filas("Auditorias")
maquinaria, _ = filas("Maquinaria")
revisiones, _ = filas("Revisiones")

print("Hoja                 Filas  Columnas")
for nombre in esperadas[1:]:
    print(f"{nombre:<20} {len(filas(nombre)[0]):>5}  {len(filas(nombre)[1])}")

zona_de = {c["ID Centro"]: c["Zona"] for c in centros}
centro_ids = set(zona_de)
emp_ids = {e["ID"] for e in empleados}
vigentes = [a for a in asignaciones if a["Fecha Fin"] is None]

for e in empleados:
    tag = e["ID"]
    alta, baja = como_fecha(e["Fecha Alta"]), como_fecha(e["Fecha Baja"])
    if e["Estado"] == "Baja":
        if not (baja and baja <= HOY):
            errores.append(f"{tag}: Baja sin fecha <= HOY")
    elif baja and baja <= HOY:
        errores.append(f"{tag}: {e['Estado']} con baja pasada")
    if e["Centro Coste"] not in centro_ids:
        errores.append(f"{tag}: centro desconocido")
    if e["Zona Asignada"] != zona_de.get(e["Centro Coste"]):
        errores.append(f"{tag}: zona no coincide con su centro")
    horas = e["Horas Semanales"]
    if e["Jornada"] == "Parcial" and horas >= 40:
        errores.append(f"{tag}: parcial con {horas}h")
    if e["Jornada"] == "Completa" and horas < 40:
        errores.append(f"{tag}: completa con {horas}h")
    if e["Salario Base"] < SMI * horas / 40 - 0.01:
        errores.append(f"{tag}: salario {e['Salario Base']} bajo SMI proporcional")
    if e["Plus Turnicidad"] and e["Turno Habitual"] != "Rotativo":
        errores.append(f"{tag}: turnicidad con turno {e['Turno Habitual']}")
    if e["Plus Nocturnidad"] and e["Turno Habitual"] != "Noche":
        errores.append(f"{tag}: nocturnidad con turno {e['Turno Habitual']}")
    if e["Estado"] != "Baja" and not any(
            a["Empleado"] == tag for a in vigentes):
        errores.append(f"{tag}: sin asignacion vigente")

if len(vigentes) < len([e for e in empleados if e["Estado"] != "Baja"]):
    errores.append("faltan asignaciones vigentes")

# contratos: todos los vigentes con fin no vencido
for c in contratos:
    ini, fin = como_fecha(c["Fecha Inicio"]), como_fecha(c["Fecha Fin"])
    emp = next((e for e in empleados if e["ID"] == c["Empleado"]), None)
    if not emp:
        errores.append(f"{c['ID']}: empleado huerfano")
        continue
    if ini > HOY:
        errores.append(f"{c['ID']}: inicio futuro")
    if fin and fin <= ini:
        errores.append(f"{c['ID']}: fin <= inicio")
    if emp["Estado"] != "Baja" and fin and fin <= HOY:
        errores.append(f"{c['ID']}: vigente vencido")
    if emp["Estado"] == "Baja" and not (fin and fin <= HOY):
        errores.append(f"{c['ID']}: baja sin fin de contrato")

# vacaciones: solapes, 30 dias, estado Vacaciones cubre HOY
por_emp = {}
for v in vacaciones:
    ini, fin = como_fecha(v["Fecha Inicio"]), como_fecha(v["Fecha Fin"])
    por_emp.setdefault(v["Empleado"], []).append((ini, fin, v["Estado"]))
    if fin < ini:
        errores.append(f"{v['ID']}: fin < inicio")
for eid, tramos in por_emp.items():
    tramos.sort()
    total = sum((f - i).days + 1 for i, f, _ in tramos)
    if total > 30:
        errores.append(f"{eid}: {total} dias > 30")
    for (i1, f1, _), (i2, _, _) in zip(tramos, tramos[1:]):
        if i2 <= f1:
            errores.append(f"{eid}: solape de vacaciones")
    emp = next((e for e in empleados if e["ID"] == eid), None)
    if emp and emp["Estado"] == "Vacaciones":
        if not any(i <= HOY <= f and e in ("Aprobada", "Disfrutada")
                   for i, f, e in tramos):
            errores.append(f"{eid}: Vacaciones sin estancia que cubre HOY")

# maquinaria y revisiones
maq_ids = {m["ID"] for m in maquinaria}
ultima = {}
for r in revisiones:
    if r["Maquinaria"] not in maq_ids:
        errores.append(f"{r['ID']}: maquinaria huerfana")
    f = como_fecha(r["Fecha"])
    if f > HOY:
        errores.append(f"{r['ID']}: revision futura")
    if not ultima.get(r["Maquinaria"]) or f > ultima[r["Maquinaria"]][0]:
        ultima[r["Maquinaria"]] = (f, r["Resultado"])
for m in maquinaria:
    prox = como_fecha(m["Proxima Revision"])
    if m["Estado"] == "Baja" and prox:
        errores.append(f"{m['ID']}: baja con revision programada")
    if m["Estado"] != "Baja":
        if not prox or prox <= HOY:
            errores.append(f"{m['ID']}: revision no futura")
        if m["Estado"] == "En mantenimiento":
            f, res = ultima.get(m["ID"], (None, None))
            if res not in ("Con hallazgos", "No conforme"):
                errores.append(f"{m['ID']}: en mantenimiento sin hallazgo")

# auditorias
for a in auditorias:
    f = como_fecha(a["Fecha"])
    if f > HOY:
        errores.append(f"{a['ID']}: auditoria futura")
    if a["Resultado"] == "Conforme" and a["Accion Correctiva"]:
        errores.append(f"{a['ID']}: conforme con accion")
    if a["Resultado"] != "Conforme" and not a["Accion Correctiva"]:
        errores.append(f"{a['ID']}: no conforme sin accion")
    fc = como_fecha(a["Fecha Cierre"])
    if fc and fc > HOY:
        errores.append(f"{a['ID']}: cierre futuro")

# proveedores
for p in proveedores:
    if p["Activo"] == "No" and p["Importe 2026"]:
        errores.append(f"{p['ID']}: inactivo con importe")
    if p["Activo"] == "Si":
        f = como_fecha(p["Ultimo Pedido"])
        if not (date(2026, 1, 1) <= f <= HOY):
            errores.append(f"{p['ID']}: activo sin pedido 2026")

# referenciales
asig_ids = {a["Empleado"] for a in vigentes}
for e in empleados:
    if e["Estado"] != "Baja" and e["ID"] not in asig_ids:
        errores.append(f"{e['ID']}: sin asignacion vigente")
prov_ids = {p["ID"] for p in proveedores}
for m in maquinaria:
    if m["Proveedor Mantenimiento"] not in prov_ids:
        errores.append(f"{m['ID']}: proveedor huerfano")
    if m["Centro"] not in centro_ids:
        errores.append(f"{m['ID']}: centro huerfano")

if errores:
    print(f"\nREAD-BACK: {len(errores)} ERRORES")
    for e in errores[:30]:
        print("  -", e)
    sys.exit(1)

print("\nREAD-BACK OK: xlsx reabierto, 0 errores de coherencia")
print(f"Empleados: {len(empleados)} | Centros: {len(centros)} | "
      f"Vigentes: {len(vigentes)} | Vacaciones: {len(vacaciones)} | "
      f"Maquinaria: {len(maquinaria)} | Revisiones: {len(revisiones)} | "
      f"Auditorias: {len(auditorias)}")
