#!/usr/bin/env python3
"""Genera data/easyexcel_demo.xlsx: dataset demo completo y coherente para EasyExcel Dashboard.

Parte de data/empleados_demo.csv (200 empleados simulados) y produce un libro
de 9 hojas con validaciones cruzadas: Empleados, Centros, Asignaciones,
Contratos, Vacaciones, Proveedores, Auditorias, Maquinaria, Revisiones.

Reglas de coherencia aplicadas (todas verificadas al final):
  - Estado "Baja" <=> Fecha Baja <= HOY.
  - Tipo de contrato separado de la jornada (Completa/Parcial); parcial < 40h.
  - Salario base >= SMI proporcional (1.221 EUR/mes 14 pagas, RD 126/2026).
  - Plus turnicidad solo turno Rotativo; nocturnidad solo turno Noche;
    penosidad solo observaciones de altura/cristaleria.
  - Zona del empleado = zona de su centro de coste.
  - Todo empleado no-baja tiene asignacion vigente en su centro de coste.
  - Vacaciones <= 30 dias naturales por empleado en 2026 y sin solapes;
    estado "Vacaciones" implica estancia que cubre HOY.
  - Proxima revision de maquinaria > HOY; ultima revision <= HOY.
  - Auditorias: acciones correctivas solo si proceden; fechas de cierre <= HOY.
"""

import csv
import math
import random
import unicodedata
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- parametros

HOY = date(2026, 10, 5)
SMI = 1221.0  # EUR/mes, 14 pagas, jornada completa (RD 126/2026, BOE 19-02-2026)
SEED = 20261005
BASE = Path(__file__).resolve().parent.parent
SRC = BASE / "data" / "empleados_demo.csv"
OUT = BASE / "data" / "easyexcel_demo.xlsx"

random.seed(SEED)

CATEGORIAS = ["Encargado", "Especialista", "Oficial", "Operario"]
TIPOS_CONTRATO = ["Indefinido", "Temporal", "Fijo Discontinuo", "Practicas"]
JORNADAS = ["Completa", "Parcial"]
TURNOS = ["Manana", "Tarde", "Noche", "Rotativo", "Partido"]
ZONAS = ["Norte", "Sur", "Este", "Oeste", "Centro", "Industrial"]
ESTADOS = ["Activo", "Vacaciones", "IT", "Suspendido", "Baja"]

BASES_40H = {
    "Encargado": (1560, 1760),
    "Especialista": (1400, 1540),
    "Oficial": (1320, 1430),
    "Operario": (1221, 1300),
}

NOMBRES_H = ["Adrian", "Alberto", "Bruno", "Carlos", "Daniel", "Eduardo", "Fernando",
             "Gonzalo", "Hector", "Ignacio", "Jorge", "Luis", "Manuel", "Miguel",
             "Pablo", "Ricardo", "Sergio", "Victor"]
NOMBRES_M = ["Alicia", "Beatriz", "Carmen", "Dolores", "Elena", "Gloria", "Ines",
             "Julia", "Laura", "Marta", "Natalia", "Patricia", "Rocio", "Silvia",
             "Teresa", "Yolanda"]
APELLIDOS = ["Alonso", "Blanco", "Carrasco", "Dominguez", "Esteban", "Fuentes",
             "Gimeno", "Herrera", "Ibanez", "Juarez", "Lozano", "Molina", "Navarro",
             "Ortega", "Prieto", "Quintana", "Ramos", "Serrano", "Tovar", "Vidal"]
LETRAS_DNI = "TRWAGMYFPDXBNJZSQVHLCKE"

OBSERVACIONES = [
    "Disponible los fines de semana",
    "Especialista en cristaleria",
    "Requiere certificacion de altura",
    "Disponible para turnos nocturnos",
    "Disponible para viajes",
    "Experiencia en hospitales",
    "En periodo de prueba",
    "",
]


# ---------------------------------------------------------------- utilidades

def sin_tildes(texto: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )


def email_de(nombre: str) -> str:
    limpio = sin_tildes(nombre).lower().replace(" ", ".")
    return f"{limpio}@empresa-limpieza.es"


def dni_nuevo() -> str:
    n = random.randint(10_000_000, 99_999_999)
    return f"{n}{LETRAS_DNI[n % 23]}"


def r5(n: float) -> int:
    """Redondea hacia arriba al multiplo de 5 (nunca por debajo del minimo)."""
    return int(math.ceil(n / 5.0) * 5)


def daterange_laborables(inicio: date, fin: date) -> int:
    dias, d = 0, inicio
    while d <= fin:
        if d.weekday() < 5:
            dias += 1
        d += timedelta(days=1)
    return dias


def add_months(d: date, meses: int) -> date:
    mes = d.month - 1 + meses
    anio = d.year + mes // 12
    mes = mes % 12 + 1
    dia = min(d.day, [31, 29 if anio % 4 == 0 and (anio % 100 != 0 or anio % 400 == 0)
                      else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][mes - 1])
    return date(anio, mes, dia)


def fecha_aleatoria(inicio: date, fin: date) -> date:
    return inicio + timedelta(days=random.randint(0, (fin - inicio).days))


def parse_fecha(txt: str):
    txt = (txt or "").strip()
    if not txt:
        return None
    a, m, d = txt.split("-")
    return date(int(a), int(m), int(d))


# ---------------------------------------------------------------- empleados

def cargar_empleados() -> list[dict]:
    with open(SRC, encoding="utf-8-sig", newline="") as fh:
        filas = list(csv.DictReader(fh))

    empleados = []
    for f in filas:
        alta = parse_fecha(f["Fecha Alta"])
        baja = parse_fecha(f["Fecha Baja"])
        estado = f["Estado"].strip()
        obs = f["Observaciones"].strip()

        # Estado <-> fecha de baja
        if baja and baja <= HOY:
            estado = "Baja"
        elif estado == "Baja" and baja is None:
            baja = max(alta + timedelta(days=5), HOY - timedelta(days=10))
            if baja > HOY:
                baja = HOY
            estado = "Baja"

        # Practicas con mas de 12 meses: el contrato vigente ya no es de practicas
        tipo = f["Contrato"].strip()
        if tipo == "Practicas":
            meses = (HOY - alta).days / 30.44
            if meses > 36:
                tipo = "Indefinido"
            elif meses > 12:
                tipo = "Temporal"
            else:
                tipo = "Practicas"

        # "Parcial" original mezclaba duracion y jornada: se resuelve abajo
        era_parcial_original = tipo == "Parcial"
        if era_parcial_original:
            tipo = "Indefinido" if random.random() < 0.5 else "Temporal"
        tipo = "Practicas" if tipo == "Practicas" else (
            "Fijo Discontinuo" if tipo == "Fijo Discontinuo" else tipo)

        horas = int(f["Horas Semanales"])
        if horas >= 40:
            jornada = "Completa"
        else:
            jornada = "Parcial"
            if horas > 35:
                horas = random.choice([30, 35])

        categoria = f["Categoría"].strip()
        if categoria == "Peón":
            categoria = "Operario"

        turno = f["Turno Habitual"].strip()
        if turno == "Rotativo":
            plus_turn = r5(random.randint(60, 120))
            plus_noct = 0
        elif turno == "Noche":
            plus_turn = 0
            plus_noct = r5(random.randint(70, 140))
        else:
            plus_turn = 0
            plus_noct = 0

        obs_low = sin_tildes(obs).lower()
        if "altura" in obs_low or "cristaleria" in obs_low:
            plus_pen = r5(random.randint(40, 100))
        else:
            plus_pen = 0

        base_lo, base_hi = BASES_40H[categoria]
        salario = r5(random.randint(base_lo, base_hi) * horas / 40)

        empleados.append({
            "id": f["ID Empleado"].strip(),
            "nombre": f["Nombre Completo"].strip(),
            "dni": f["DNI/NIE"].strip(),
            "telefono": f["Teléfono"].strip(),
            "email": email_de(f["Nombre Completo"].strip()),
            "categoria": categoria,
            "tipo_contrato": tipo,
            "jornada": jornada,
            "alta": alta,
            "baja": baja,
            "centro": f["Centro Coste"].strip().replace("CC", "CC0") if len(f["Centro Coste"].strip()) == 3 else f["Centro Coste"].strip(),
            "turno": turno,
            "horas": horas,
            "salario": salario,
            "plus_turn": plus_turn,
            "plus_noct": plus_noct,
            "plus_pen": plus_pen,
            "estado": estado,
            "observaciones": obs,
        })

    # normaliza IDs de centro: CC1 -> CC01
    for e in empleados:
        c = e["centro"]
        if c.startswith("CC") and c[2:].isdigit() and len(c) == 3:
            e["centro"] = "CC0" + c[2:]

    return empleados


def generar_nuevos(empleados: list[dict]) -> list[dict]:
    """15 empleados nuevos para los centros CC06 y CC07."""
    nuevos = []
    distribucion = (["CC06"] * 8) + (["CC07"] * 7)
    categorias_nuevas = ["Encargado", "Encargado", "Oficial", "Oficial", "Especialista",
                         "Especialista", "Especialista", "Operario", "Operario"]
    random.shuffle(categorias_nuevas)

    for i, centro in enumerate(distribucion, start=201):
        nombre = f"{random.choice(NOMBRES_H if i % 2 else NOMBRES_M)} {random.choice(APELLIDOS)} {random.choice(APELLIDOS)}"
        categoria = categorias_nuevas[i - 201] if i - 201 < len(categorias_nuevas) else "Operario"
        turno = random.choice(TURNOS)
        horas = random.choice([20, 25, 30, 35, 40])
        jornada = "Completa" if horas == 40 else "Parcial"
        tipo = random.choice(["Indefinido", "Indefinido", "Temporal", "Fijo Discontinuo"])
        alta = fecha_aleatoria(date(2023, 1, 1), date(2026, 9, 1))
        obs = random.choice(OBSERVACIONES)

        if turno == "Rotativo":
            plus_turn, plus_noct = r5(random.randint(60, 120)), 0
        elif turno == "Noche":
            plus_turn, plus_noct = 0, r5(random.randint(70, 140))
        else:
            plus_turn, plus_noct = 0, 0
        obs_low = sin_tildes(obs).lower()
        plus_pen = r5(random.randint(40, 100)) if ("altura" in obs_low or "cristaleria" in obs_low) else 0

        base_lo, base_hi = BASES_40H[categoria]
        estado = random.choices(ESTADOS, weights=[84, 5, 4, 5, 2])[0]
        baja = None
        if estado == "Baja":
            estado_baja = fecha_aleatoria(max(alta + timedelta(days=5), date(2026, 1, 5)), HOY - timedelta(days=2))
            if estado_baja > alta:
                baja = estado_baja

        nuevos.append({
            "id": f"EMP{i}",
            "nombre": nombre,
            "dni": dni_nuevo(),
            "telefono": f"6{random.randint(10000000, 99999999)}",
            "email": email_de(nombre),
            "categoria": categoria,
            "tipo_contrato": tipo,
            "jornada": jornada,
            "alta": alta,
            "baja": baja,
            "centro": centro,
            "turno": turno,
            "horas": horas,
            "salario": r5(random.randint(base_lo, base_hi) * horas / 40),
            "plus_turn": plus_turn,
            "plus_noct": plus_noct,
            "plus_pen": plus_pen,
            "estado": estado,
            "observaciones": obs,
        })

    empleados.extend(nuevos)
    return empleados


# ---------------------------------------------------------------- hojas

def construir_asignaciones(empleados: list[dict]) -> list[dict]:
    asignaciones = []
    contador = 0
    for e in empleados:
        centro_actual = e["centro"]
        if e["estado"] == "Baja":
            asignaciones.append({
                "id": f"ASG{contador:03d}", "empleado": e["id"], "centro": centro_actual,
                "puesto": e["categoria"], "turno": e["turno"], "horas": e["horas"],
                "inicio": e["alta"], "fin": e["baja"] or HOY - timedelta(days=1),
            })
            contador += 1
            continue

        rota = random.random() < 0.30 and e["alta"] <= date(2025, 6, 1)
        if rota:
            otros = [c for c in CENTROS_IDS if c != centro_actual]
            centro_hist = random.choice(otros)
            inicio_vig = fecha_aleatoria(max(e["alta"], date(2025, 1, 1)), HOY - timedelta(days=7))
            fin_hist = inicio_vig - timedelta(days=1)
            inicio_hist = max(e["alta"], fin_hist - timedelta(days=random.randint(300, 700)))
            if inicio_hist < fin_hist:
                asignaciones.append({
                    "id": f"ASG{contador:03d}", "empleado": e["id"], "centro": centro_hist,
                    "puesto": e["categoria"], "turno": e["turno"], "horas": e["horas"],
                    "inicio": inicio_hist, "fin": fin_hist,
                })
                contador += 1
            inicio_vig_real = inicio_vig
        else:
            inicio_vig_real = e["alta"]

        asignaciones.append({
            "id": f"ASG{contador:03d}", "empleado": e["id"], "centro": centro_actual,
            "puesto": e["categoria"], "turno": e["turno"], "horas": e["horas"],
            "inicio": inicio_vig_real, "fin": None,
        })
        contador += 1
    return asignaciones


def construir_contratos(empleados: list[dict]) -> list[dict]:
    contratos = []
    for i, e in enumerate(sorted(empleados, key=lambda x: x["id"]), start=1):
        tipo = e["tipo_contrato"]
        if e["estado"] == "Baja":
            inicio = e["alta"]
            fin = e["baja"]
            if tipo == "Indefinido" and fin and (fin - inicio).days < 90:
                pass
        elif tipo == "Temporal":
            inicio = max(e["alta"], HOY - timedelta(days=random.randint(120, 420)))
            fin = add_months(inicio, random.randint(6, 12))
            if fin <= HOY:
                fin = add_months(HOY, random.randint(2, 8))
        elif tipo == "Practicas":
            inicio = e["alta"]
            fin = add_months(inicio, random.choice([6, 12]))
            if fin <= HOY:
                fin = add_months(HOY, 3)
        else:  # Indefinido / Fijo Discontinuo
            inicio = e["alta"]
            fin = None

        duracion = None
        if fin:
            duracion = max(1, round((fin - inicio).days / 30.44))

        contratos.append({
            "id": f"CTR{i:03d}",
            "empleado": e["id"],
            "tipo": tipo,
            "jornada": e["jornada"],
            "horas": e["horas"],
            "inicio": inicio,
            "fin": fin,
            "duracion_meses": duracion,
            "prorrogas": random.randint(0, 3) if tipo == "Temporal" else 0,
            "documento": f"/contratos/{e['id']}_{inicio.year}.pdf",
        })
    return contratos


def construir_vacaciones(empleados: list[dict]) -> list[dict]:
    vacaciones = []
    contador = 0

    def add(e, inicio, fin, estado):
        nonlocal contador
        if fin < inicio:
            return
        nat = (fin - inicio).days + 1
        vacaciones.append({
            "id": f"VAC{contador:03d}", "empleado": e["id"], "anio": 2026,
            "inicio": inicio, "fin": fin, "naturales": nat,
            "laborables": daterange_laborables(inicio, fin),
            "solicitud": inicio - timedelta(days=random.randint(20, 60)),
            "estado": estado,
        })
        contador += 1

    for e in empleados:
        if e["estado"] == "Baja":
            continue
        usado = 0

        if e["estado"] == "Vacaciones":
            d = random.randint(7, 10)
            ini = HOY - timedelta(days=random.randint(1, 6))
            add(e, ini, ini + timedelta(days=d - 1), "Disfrutada")
            usado += d

        if e["alta"] <= date(2026, 6, 1):
            d = min(random.randint(7, 14), 30 - usado)
            if d >= 3:
                ini = fecha_aleatoria(date(2026, 6, 15), date(2026, 7, 25))
                add(e, ini, ini + timedelta(days=d - 1), "Disfrutada")
                usado += d

        if random.random() < 0.50 and usado <= 20:
            d = min(random.randint(5, 10), 30 - usado)
            if d >= 3:
                ini = fecha_aleatoria(date(2026, 11, 5), date(2026, 11, 30))
                add(e, ini, ini + timedelta(days=d - 1), "Aprobada")
                usado += d

        if random.random() < 0.30 and usado <= 24:
            d = min(random.randint(3, 6), 30 - usado)
            if d >= 3:
                ini = fecha_aleatoria(date(2026, 12, 15), date(2026, 12, 26))
                add(e, ini, ini + timedelta(days=d - 1), "Pendiente")
                usado += d

    return vacaciones


CENTROS_DEF = [
    ("CC01", "Hospital General San Vicente", "Servicios Sanitarios San Vicente",
     "Av. dels Servents 120, Valencia", "Sur", "963110201", date(2019, 9, 1), 12400),
    ("CC02", "Torre Empresarial Turia", "Gestion Inmobiliaria Turia",
     "Gran Via Marques del Turia 100, Valencia", "Centro", "963110202", date(2020, 3, 16), 4500),
    ("CC03", "Centro Comercial Aqua Plaza", "Aqua Plaza Gestion",
     "Av. del Port 12, Valencia", "Oeste", "963110203", date(2021, 5, 3), 18200),
    ("CC04", "Oficinas Rio Verde", "Rio Verde Consulting",
     "Camí de la Mar 8, Alboraya", "Este", "963110204", date(2022, 1, 10), 2200),
    ("CC05", "Poligon Industrial Quart", "Grupo Industrial Quart",
     "Av. de l'Enginyer Llorens 7, Quart de Poblet", "Industrial", "963110205", date(2023, 7, 1), 6100),
    ("CC06", "Residencial Benimaclet", "Comunidad Benimaclet 44",
     "Carrer de Denia 55, Valencia", "Norte", "963110206", date(2024, 2, 1), 3500),
    ("CC07", "Mercado Central Nou", "Mercado Central Nou SL",
     "Plaça del Mercat Central 3, Valencia", "Centro", "963110207", date(2025, 9, 15), 5200),
]
CENTROS_IDS = [c[0] for c in CENTROS_DEF]


def construir_centros(empleados: list[dict], asignaciones: list[dict]) -> list[dict]:
    vigentes = [a for a in asignaciones if a["fin"] is None]
    centros = []
    for cid, nombre, cliente, dir_, zona, tel, inicio, m2 in CENTROS_DEF:
        plantilla = [a for a in vigentes if a["centro"] == cid]
        candidatos = [
            e for e in empleados
            if e["centro"] == cid and e["estado"] in ("Activo", "Vacaciones")
        ]
        responsables = [e for e in candidatos if e["categoria"] == "Encargado"]
        pool = responsables or candidatos
        resp = min(pool, key=lambda e: e["alta"])["id"] if pool else ""
        centros.append({
            "id": cid, "nombre": nombre, "cliente": cliente, "direccion": dir_,
            "zona": zona, "telefono": tel, "responsable": resp,
            "inicio_servicio": inicio, "superficie_m2": m2,
            "plantilla_asignada": len(plantilla), "estado": "Activo",
        })
    return centros


def construir_proveedores() -> list[dict]:
    defs = [
        ("Suministros Levante SL", "Material de limpieza"),
        ("HigienePro Valencia SL", "Papel y higiene"),
        ("QuimValencia SA", "Productos quimicos"),
        ("Maquinaria CleanTech SL", "Maquinaria"),
        ("Mantenimientos Integrales Edificios SL", "Mantenimiento de instalaciones"),
        ("Ascensores Turia SL", "Mantenimiento de ascensores"),
        ("EPI Seguridad Levante SL", "Seguridad y EPIs"),
        ("Jardines del Turia SL", "Jardineria"),
        ("Gestion Residuos Mediterraneo SL", "Gestion de residuos"),
        ("Lavanderia Blanca Nieve SL", "Lavanderia industrial"),
        ("Cristal Levante SL", "Cristaleria y altura"),
        ("Transportes Ronda SL", "Transporte y logistica"),
        ("Papeleria Oficina Turia SL", "Material de oficina"),
        ("Control de Plagas Levante SL", "Desinsectizacion y desratizacion"),
        ("Elevadores del Este SL", "Mantenimiento de plataformas"),
        ("Aguas Mediterraneo SL", "Servicio de agua"),
        ("Electricidad Edificios Sur SL", "Mantenimiento electrico"),
        ("Climatizacion Valencia Norte SL", "Mantenimiento HVAC"),
    ]
    proveedores = []
    for i, (nombre, servicio) in enumerate(defs, start=1):
        activo = i not in (13, 16)
        if activo:
            ultimo = fecha_aleatoria(date(2026, 1, 5), HOY - timedelta(days=3))
            importe = random.randint(50, 4500) * 10
        else:
            ultimo = fecha_aleatoria(date(2024, 3, 1), date(2025, 11, 30))
            importe = 0
        dom = sin_tildes(nombre.lower().split(" sl")[0].split(" sa")[0]).replace(" ", "")
        proveedores.append({
            "id": f"PROV{i:02d}",
            "nombre": nombre,
            "cif": f"B{random.randint(10000000, 99999999)}",
            "servicio": servicio,
            "contacto": f"{random.choice(NOMBRES_M)} {random.choice(APELLIDOS)}",
            "telefono": f"96{random.randint(1000000, 9999999)}",
            "email": f"contacto@{dom}.es",
            "direccion": f"Carrer de {random.choice(APELLIDOS)} {random.randint(2, 99)}, Valencia",
            "activo": "Si" if activo else "No",
            "ultimo_pedido": ultimo,
            "importe_2026": importe,
            "puntuacion": random.randint(3, 5),
        })
    return proveedores


def construir_maquinaria(centros: list[dict], proveedores: list[dict]) -> list[dict]:
    tipos = [
        ("Fregadora discos", "Nilfisk", "Sc6000", 2400),
        ("Fregadora autonoma", "Karcher", "B 90 R", 6800),
        ("Aspiradora mochila", "Numatic", "Henry XWD", 380),
        ("Aspiradora industrial", "Nilfisk", "VA2000", 1250),
        ("Monobrocha", "Viper", "Fang 150", 560),
        ("Hidrolavadora", "Karcher", "HD 6/15", 990),
        ("Secador de suelos", "Numatic", "TMR 390", 1450),
        ("Elevador tijera", "JLG", "1930ES", 7500),
        ("Carro de limpieza", "Vikan", "Compact 250", 320),
        ("Pulidora orbital", "3M", "FriXion 5200", 1100),
    ]
    proveedores_mant = [p for p in proveedores
                        if p["servicio"] in ("Maquinaria", "Mantenimiento de plataformas",
                                             "Mantenimiento de instalaciones")]
    maquinaria = []
    n = 0
    for c in centros:
        cantidad = 8 if c["id"] in ("CC01", "CC03") else 7
        for _ in range(cantidad):
            n += 1
            tipo, marca, modelo, coste = random.choice(tipos)
            compra = fecha_aleatoria(date(2018, 1, 1), date(2025, 6, 30))
            garantia = add_months(compra, 24)
            estado = random.choices(["Operativa", "En mantenimiento", "Baja"],
                                    weights=[80, 13, 7])[0]
            proxima = None if estado == "Baja" else HOY + timedelta(days=random.randint(5, 90))
            maquinaria.append({
                "id": f"MAQ{n:03d}",
                "tipo": tipo,
                "marca": marca,
                "modelo": modelo,
                "serie": f"{marca[:2].upper()}{random.randint(100000, 999999)}",
                "centro": c["id"],
                "fecha_compra": compra,
                "garantia_hasta": garantia,
                "coste": coste,
                "estado": estado,
                "proxima_revision": proxima,
                "proveedor_mantenimiento": random.choice(proveedores_mant)["id"],
            })
    return maquinaria


def construir_revisiones(maquinaria: list[dict]) -> list[dict]:
    revisiones = []
    n = 0
    for m in maquinaria:
        if m["estado"] == "Baja":
            continue
        num = random.randint(1, 3)
        fechas = sorted(fecha_aleatoria(date(2026, 1, 5), HOY - timedelta(days=3))
                        for _ in range(num))
        for k, fecha in enumerate(fechas):
            n += 1
            es_ultima = k == num - 1
            if es_ultima and m["estado"] == "En mantenimiento":
                tipo, resultado = "Correctiva", "Con hallazgos"
            elif es_ultima:
                tipo = random.choice(["Preventiva", "Preventiva", "ITV"])
                resultado = random.choices(["Conforme", "Con hallazgos"],
                                           weights=[85, 15])[0]
            else:
                tipo = random.choice(["Preventiva", "Correctiva", "ITV"])
                resultado = random.choices(["Conforme", "Con hallazgos", "No conforme"],
                                           weights=[75, 20, 5])[0]
            revisiones.append({
                "id": f"REV{n:03d}",
                "maquinaria": m["id"],
                "fecha": fecha,
                "tipo": tipo,
                "resultado": resultado,
                "tecnico": f"{random.choice(NOMBRES_H)} {random.choice(APELLIDOS)}",
                "coste": 0 if tipo == "Preventiva" else random.randint(5, 35) * 10,
                "observaciones": {
                    "Conforme": "Equipo operativo, sin observaciones",
                    "Con hallazgos": "Se sustituyen filtros y se revisa instalacion electrica",
                    "No conforme": "Equipo parado hasta sustitucion de componente principal",
                }[resultado],
            })
    return revisiones


def construir_auditorias(centros: list[dict]) -> list[dict]:
    hallazgos = [
        "Restos en zona de carga y descarga del sotano",
        "Producto quimico sin ficha de seguridad actualizada",
        "Falta de EPIs en el armario del centro",
        "Programa de limpieza sin firma del responsable",
        "Contenedor de residuos sin etiquetar",
        "Zona comun con incidencia de limpieza recurrente",
        "Maquinaria sin etiqueta de proxima revision visible",
        "Deposito de agua estancado en patio interior",
    ]
    acciones = [
        "Limpieza profunda programada y ficha SDS actualizada",
        "Reposicion de EPIs y formacion al equipo del turno",
        "Firma y archivo del programa de limpieza",
        "Contenedor etiquetado y formacion de segregacion",
        "Revision semanal anadida al plan del centro",
        "Revision preventiva adelantada y etiquetado",
    ]
    auditores = ["Elena Sanchis", "Marco Ferrandis", "Laura Bouchar",
                 "Sergio Picó", "Nerea Tortajada"]
    auditorias = []
    n = 0
    for c in centros:
        for _ in range(random.randint(4, 6)):
            n += 1
            fecha = fecha_aleatoria(date(2026, 1, 10), HOY - timedelta(days=5))
            puntuacion = random.randint(72, 99)
            if puntuacion >= 90:
                resultado, hallazgo, accion = "Conforme", "Sin hallazgos relevantes", ""
            elif puntuacion >= 80:
                resultado = "Acciones correctivas"
                hallazgo = random.choice(hallazgos)
                accion = random.choice(acciones)
            else:
                resultado = "No conforme"
                hallazgo = random.choice(hallazgos)
                accion = random.choice(acciones)
            estado = "Cerrada"
            cierre = fecha + timedelta(days=random.randint(7, 30))
            if accion and cierre > HOY:
                estado = "Abierta"
                cierre = None
            elif not accion:
                cierre = None
            auditorias.append({
                "id": f"AUD{n:03d}",
                "centro": c["id"],
                "fecha": fecha,
                "tipo": random.choice(["Interna", "Cliente", "Externa"]),
                "inspector": random.choice(auditores),
                "puntuacion": puntuacion,
                "resultado": resultado,
                "hallazgos": hallazgo,
                "accion": accion,
                "fecha_cierre": cierre,
                "estado": estado,
            })
    return auditorias


# ---------------------------------------------------------------- validacion

def validar(empleados, centros, asignaciones, contratos, vacaciones,
            proveedores, auditorias, maquinaria, revisiones) -> list[str]:
    errores = []
    ids_emp = {e["id"] for e in empleados}
    ids_centro = {c["id"] for c in centros}
    zona_de = {c["id"]: c["zona"] for c in centros}
    vigentes = [a for a in asignaciones if a["fin"] is None]
    min_salario = SMI

    for e in empleados:
        tag = e["id"]
        if e["estado"] == "Baja":
            if not (e["baja"] and e["baja"] <= HOY):
                errores.append(f"{tag}: estado Baja sin fecha de baja <= HOY")
        else:
            if e["baja"] and e["baja"] <= HOY:
                errores.append(f"{tag}: estado {e['estado']} con fecha de baja pasada")
        if e["centro"] not in ids_centro:
            errores.append(f"{tag}: centro {e['centro']} inexistente")
        if e["estado"] != "Baja" and zona_de.get(e["centro"]) and \
                e.get("zona") and e["zona"] != zona_de[e["centro"]]:
            errores.append(f"{tag}: zona {e['zona']} != zona del centro")
        if e["jornada"] == "Parcial" and e["horas"] >= 40:
            errores.append(f"{tag}: jornada Parcial con {e['horas']}h")
        if e["jornada"] == "Completa" and e["horas"] < 40:
            errores.append(f"{tag}: jornada Completa con {e['horas']}h")
        if e["salario"] < min_salario * e["horas"] / 40 - 0.01:
            errores.append(f"{tag}: salario {e['salario']} < SMI proporcional "
                           f"({min_salario * e['horas'] / 40:.2f})")
        if e["plus_turn"] and e["turno"] != "Rotativo":
            errores.append(f"{tag}: plus turnicidad con turno {e['turno']}")
        if e["plus_noct"] and e["turno"] != "Noche":
            errores.append(f"{tag}: plus nocturnidad con turno {e['turno']}")
        obs_low = sin_tildes(e["observaciones"]).lower()
        if e["plus_pen"] and not ("altura" in obs_low or "cristaleria" in obs_low):
            errores.append(f"{tag}: plus penosidad sin motivo")
        if not e["plus_pen"] and ("altura" in obs_low or "cristaleria" in obs_low):
            errores.append(f"{tag}: motivo de penosidad sin plus")
        if " " in e["email"] or sin_tildes(e["email"]) != e["email"].lower():
            errores.append(f"{tag}: email no normalizado: {e['email']}")

    for a in asignaciones:
        if a["empleado"] not in ids_emp:
            errores.append(f"{a['id']}: empleado {a['empleado']} inexistente")
        if a["centro"] not in ids_centro:
            errores.append(f"{a['id']}: centro {a['centro']} inexistente")
        if a["fin"] and a["fin"] < a["inicio"]:
            errores.append(f"{a['id']}: fin < inicio")

    emp_por_id = {e["id"]: e for e in empleados}
    for e in empleados:
        if e["estado"] == "Baja":
            continue
        mios = [a for a in vigentes if a["empleado"] == e["id"]]
        if not mios:
            errores.append(f"{e['id']}: sin asignacion vigente")
        elif mios[0]["centro"] != e["centro"]:
            errores.append(f"{e['id']}: asignacion vigente en {mios[0]['centro']} "
                           f"!= centro {e['centro']}")

    for c in contratos:
        e = emp_por_id.get(c["empleado"])
        if not e:
            errores.append(f"{c['id']}: empleado inexistente")
            continue
        if c["fin"] and c["fin"] <= c["inicio"]:
            errores.append(f"{c['id']}: fin <= inicio")
        if c["fin"] and e["estado"] != "Baja" and c["fin"] <= HOY:
            errores.append(f"{c['id']}: contrato vigente ya vencido")
        if e["estado"] == "Baja" and not (c["fin"] and c["fin"] <= HOY):
            errores.append(f"{c['id']}: empleado de baja sin fecha fin")
        if c["tipo"] == "Practicas" and (HOY - c["inicio"]).days > 400:
            errores.append(f"{c['id']}: practicas > 12 meses")
        if c["jornada"] == "Parcial" and c["horas"] >= 40:
            errores.append(f"{c['id']}: contrato parcial a {c['horas']}h")

    vac_emp = {}
    for v in vacaciones:
        vac_emp.setdefault(v["empleado"], []).append(v)
        if v["empleado"] not in ids_emp:
            errores.append(f"{v['id']}: empleado inexistente")
        if v["fin"] < v["inicio"]:
            errores.append(f"{v['id']}: fin < inicio")
    for eid, lista in vac_emp.items():
        lista.sort(key=lambda v: v["inicio"])
        total = sum(v["naturales"] for v in lista)
        if total > 30:
            errores.append(f"{eid}: {total} dias naturales de vacaciones > 30")
        for a, b in zip(lista, lista[1:]):
            if b["inicio"] <= a["fin"]:
                errores.append(f"{eid}: vacaciones solapadas ({a['id']}/{b['id']})")
        e = emp_por_id.get(eid)
        if e and e["estado"] == "Vacaciones":
            if not any(v["inicio"] <= HOY <= v["fin"] and
                       v["estado"] in ("Aprobada", "Disfrutada") for v in lista):
                errores.append(f"{eid}: estado Vacaciones sin estancia que cubra HOY")

    prov_ids = {p["id"] for p in proveedores}
    for p in proveedores:
        if p["activo"] == "No" and p["importe_2026"] != 0:
            errores.append(f"{p['id']}: proveedor inactivo con importe 2026")
        if p["activo"] == "Si" and not (date(2026, 1, 1) <= p["ultimo_pedido"] <= HOY):
            errores.append(f"{p['id']}: proveedor activo sin pedido en 2026")

    maq_ids = {m["id"] for m in maquinaria}
    for m in maquinaria:
        if m["centro"] not in ids_centro:
            errores.append(f"{m['id']}: centro inexistente")
        if m["proveedor_mantenimiento"] not in prov_ids:
            errores.append(f"{m['id']}: proveedor mantenimiento inexistente")
        if m["estado"] != "Baja":
            if not m["proxima_revision"] or m["proxima_revision"] <= HOY:
                errores.append(f"{m['id']}: proxima revision no futura")
        elif m["proxima_revision"]:
            errores.append(f"{m['id']}: maquinaria de baja con proxima revision")

    rev_ultima = {}
    for r in revisiones:
        if r["maquinaria"] not in maq_ids:
            errores.append(f"{r['id']}: maquinaria inexistente")
        if r["fecha"] > HOY:
            errores.append(f"{r['id']}: revision futura")
        m = rev_ultima.get(r["maquinaria"])
        if not m or r["fecha"] > m["fecha"]:
            rev_ultima[r["maquinaria"]] = r
    maq_por_id = {m["id"]: m for m in maquinaria}
    for mid, r in rev_ultima.items():
        m = maq_por_id[mid]
        if m["estado"] == "En mantenimiento" and r["resultado"] not in (
                "Con hallazgos", "No conforme"):
            errores.append(f"{mid}: en mantenimiento con ultima revision {r['resultado']}")

    for a in auditorias:
        if a["centro"] not in ids_centro:
            errores.append(f"{a['id']}: centro inexistente")
        if a["fecha"] > HOY:
            errores.append(f"{a['id']}: auditoria futura")
        if a["resultado"] == "Conforme" and a["accion"]:
            errores.append(f"{a['id']}: conforme con accion correctiva")
        if a["resultado"] != "Conforme" and not a["accion"]:
            errores.append(f"{a['id']}: no conforme sin accion correctiva")
        if a["fecha_cierre"]:
            if a["fecha_cierre"] > HOY:
                errores.append(f"{a['id']}: cierre futuro")
            if a["estado"] != "Cerrada":
                errores.append(f"{a['id']}: con cierre pero estado {a['estado']}")
        elif a["estado"] == "Cerrada" and a["accion"]:
            errores.append(f"{a['id']}: accion pendiente marcada Cerrada")

    return errores


# ---------------------------------------------------------------- escritura

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
HEADER_FONT = Font(bold=True)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def hoja(wb, nombre, columnas, filas, formatos=None, primero=False):
    ws = wb.active if primero else wb.create_sheet()
    ws.title = nombre
    formatos = formatos or {}
    ws.append(columnas)
    for celda in ws[1]:
        celda.fill = HEADER_FILL
        celda.font = HEADER_FONT
        celda.border = BORDER
        celda.alignment = Alignment(vertical="center")
    for fila in filas:
        ws.append(fila)
    for idx, col in enumerate(columnas, start=1):
        fmt = formatos.get(col)
        for row in range(2, ws.max_row + 1):
            celda = ws.cell(row=row, column=idx)
            celda.border = BORDER
            if fmt:
                celda.number_format = fmt
        largo = max([len(str(col))] +
                    [len(str(ws.cell(row=r, column=idx).value or ""))
                     for r in range(2, min(ws.max_row, 60) + 1)])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(largo + 2, 10), 42)
    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = ws.dimensions
    return ws


def escribir(libro_data: dict, info: list[str]) -> None:
    wb = Workbook()
    fmt_fecha = "DD/MM/YYYY"
    fmt_eur = '#,##0.00 "€"'

    # Info
    ws = wb.active
    ws.title = "Info"
    ws.column_dimensions["A"].width = 110
    for i, linea in enumerate(info, start=1):
        celda = ws.cell(row=i, column=1, value=linea)
        if i == 1 or linea.endswith(":"):
            celda.font = Font(bold=True, size=12 if i == 1 else 11)
        celda.alignment = Alignment(wrap_text=False)

    e = libro_data["empleados"]
    hoja(wb, "Empleados",
         ["ID", "Nombre Completo", "DNI/NIE", "Telefono", "Email", "Categoria",
          "Tipo Contrato", "Jornada", "Fecha Alta", "Fecha Baja", "Centro Coste",
          "Zona Asignada", "Turno Habitual", "Horas Semanales", "Salario Base",
          "Plus Turnicidad", "Plus Nocturnidad", "Plus Penosidad", "Estado",
          "Observaciones", "Foto URL"],
         [[x["id"], x["nombre"], x["dni"], x["telefono"], x["email"], x["categoria"],
           x["tipo_contrato"], x["jornada"], x["alta"], x["baja"], x["centro"],
           x["zona"], x["turno"], x["horas"], x["salario"], x["plus_turn"],
           x["plus_noct"], x["plus_pen"], x["estado"], x["observaciones"],
           f"https://demo.easyexcel.local/fotos/{x['id']}.jpg"] for x in e],
         formatos={"Fecha Alta": fmt_fecha, "Fecha Baja": fmt_fecha,
                   "Salario Base": fmt_eur, "Plus Turnicidad": fmt_eur,
                   "Plus Nocturnidad": fmt_eur, "Plus Penosidad": fmt_eur})

    c = libro_data["centros"]
    hoja(wb, "Centros",
         ["ID Centro", "Nombre", "Cliente", "Direccion", "Zona", "Telefono",
          "Responsable", "Inicio Servicio", "Superficie m2", "Plantilla Asignada",
          "Estado"],
         [[x["id"], x["nombre"], x["cliente"], x["direccion"], x["zona"],
           x["telefono"], x["responsable"], x["inicio_servicio"],
           x["superficie_m2"], x["plantilla_asignada"], x["estado"]] for x in c],
         formatos={"Inicio Servicio": fmt_fecha})

    a = libro_data["asignaciones"]
    hoja(wb, "Asignaciones",
         ["ID", "Empleado", "Centro", "Puesto", "Turno", "Horas Semanales",
          "Fecha Inicio", "Fecha Fin"],
         [[x["id"], x["empleado"], x["centro"], x["puesto"], x["turno"],
           x["horas"], x["inicio"], x["fin"]] for x in a],
         formatos={"Fecha Inicio": fmt_fecha, "Fecha Fin": fmt_fecha})

    ct = libro_data["contratos"]
    hoja(wb, "Contratos",
         ["ID", "Empleado", "Tipo", "Jornada", "Horas Semanales", "Fecha Inicio",
          "Fecha Fin", "Duracion Meses", "Prorrogas", "Documento"],
         [[x["id"], x["empleado"], x["tipo"], x["jornada"], x["horas"],
           x["inicio"], x["fin"], x["duracion_meses"], x["prorrogas"],
           x["documento"]] for x in ct],
         formatos={"Fecha Inicio": fmt_fecha, "Fecha Fin": fmt_fecha})

    v = libro_data["vacaciones"]
    hoja(wb, "Vacaciones",
         ["ID", "Empleado", "Anio", "Fecha Inicio", "Fecha Fin", "Dias Naturales",
          "Dias Laborables", "Fecha Solicitud", "Estado"],
         [[x["id"], x["empleado"], x["anio"], x["inicio"], x["fin"],
           x["naturales"], x["laborables"], x["solicitud"], x["estado"]]
          for x in v],
         formatos={"Fecha Inicio": fmt_fecha, "Fecha Fin": fmt_fecha,
                   "Fecha Solicitud": fmt_fecha})

    p = libro_data["proveedores"]
    hoja(wb, "Proveedores",
         ["ID", "Nombre", "CIF", "Servicio", "Contacto", "Telefono", "Email",
          "Direccion", "Activo", "Ultimo Pedido", "Importe 2026", "Puntuacion"],
         [[x["id"], x["nombre"], x["cif"], x["servicio"], x["contacto"],
           x["telefono"], x["email"], x["direccion"], x["activo"],
           x["ultimo_pedido"], x["importe_2026"], x["puntuacion"]] for x in p],
         formatos={"Ultimo Pedido": fmt_fecha, "Importe 2026": fmt_eur})

    au = libro_data["auditorias"]
    hoja(wb, "Auditorias",
         ["ID", "Centro", "Fecha", "Tipo", "Inspector", "Puntuacion", "Resultado",
          "Hallazgos", "Accion Correctiva", "Fecha Cierre", "Estado"],
         [[x["id"], x["centro"], x["fecha"], x["tipo"], x["inspector"],
           x["puntuacion"], x["resultado"], x["hallazgos"], x["accion"],
           x["fecha_cierre"], x["estado"]] for x in au],
         formatos={"Fecha": fmt_fecha, "Fecha Cierre": fmt_fecha})

    m = libro_data["maquinaria"]
    hoja(wb, "Maquinaria",
         ["ID", "Tipo", "Marca", "Modelo", "N Serie", "Centro", "Fecha Compra",
          "Garantia Hasta", "Coste", "Estado", "Proxima Revision",
          "Proveedor Mantenimiento"],
         [[x["id"], x["tipo"], x["marca"], x["modelo"], x["serie"], x["centro"],
           x["fecha_compra"], x["garantia_hasta"], x["coste"], x["estado"],
           x["proxima_revision"], x["proveedor_mantenimiento"]] for x in m],
         formatos={"Fecha Compra": fmt_fecha, "Garantia Hasta": fmt_fecha,
                   "Proxima Revision": fmt_fecha, "Coste": fmt_eur})

    r = libro_data["revisiones"]
    hoja(wb, "Revisiones",
         ["ID", "Maquinaria", "Fecha", "Tipo", "Resultado", "Tecnico", "Coste",
          "Observaciones"],
         [[x["id"], x["maquinaria"], x["fecha"], x["tipo"], x["resultado"],
           x["tecnico"], x["coste"], x["observaciones"]] for x in r],
         formatos={"Fecha": fmt_fecha, "Coste": fmt_eur})

    wb.save(OUT)


# ---------------------------------------------------------------- main

def main() -> int:
    empleados = cargar_empleados()
    generar_nuevos(empleados)

    # zona heredada del centro
    zona_de = {c[0]: c[4] for c in CENTROS_DEF}
    for e in empleados:
        e["zona"] = zona_de[e["centro"]]

    asignaciones = construir_asignaciones(empleados)
    contratos = construir_contratos(empleados)
    vacaciones = construir_vacaciones(empleados)
    centros = construir_centros(empleados, asignaciones)
    proveedores = construir_proveedores()
    maquinaria = construir_maquinaria(centros, proveedores)
    revisiones = construir_revisiones(maquinaria)
    auditorias = construir_auditorias(centros)

    datos = {
        "empleados": empleados, "centros": centros,
        "asignaciones": asignaciones, "contratos": contratos,
        "vacaciones": vacaciones, "proveedores": proveedores,
        "auditorias": auditorias, "maquinaria": maquinaria,
        "revisiones": revisiones,
    }

    errores = validar(**datos)
    if errores:
        print(f"VALIDACION: {len(errores)} errores")
        for err in errores[:40]:
            print("  -", err)
        return 1

    info = [
        "EasyExcel Dashboard - Dataset de demostracion",
        "",
        "Datos 100% sinteticos. Ningun dato real de ISS ni de personas reales.",
        "Generado por scripts/generar_dataset_demo.py desde data/empleados_demo.csv (base de Jorge).",
        f"Fecha de generacion: {HOY.strftime('%d/%m/%Y')}.",
        "",
        "Hojas:",
        f"  - Empleados: {len(empleados)} registros (200 base + 15 nuevos en CC06/CC07).",
        f"  - Centros: {len(centros)} centros de Valencia con cliente, zona y plantilla.",
        f"  - Asignaciones: {len(asignaciones)} historiales empleado-centro con rotaciones.",
        f"  - Contratos: {len(contratos)} contratos vigentes con tipo, jornada y vencimiento.",
        f"  - Vacaciones: {len(vacaciones)} peticiones 2026 (max 30 dias naturales/persona).",
        f"  - Proveedores: {len(proveedores)} proveedores por servicio.",
        f"  - Auditorias: {len(auditorias)} auditorias con puntuacion y acciones correctivas.",
        f"  - Maquinaria: {len(maquinaria)} equipos con garantia y proxima revision.",
        f"  - Revisiones: {len(revisiones)} revisiones de maquinaria en 2026.",
        "",
        "Reglas de coherencia aplicadas:",
        f"  - Estado 'Baja' <=> Fecha Baja <= {HOY.strftime('%d/%m/%Y')}.",
        "  - Tipo de contrato (duración) separado de Jornada (Completa/Parcial); parcial < 40h.",
        f"  - Salario base >= SMI proporcional: {SMI:.0f} EUR/mes 14 pagas a 40h",
        "    (Real Decreto 126/2026, BOE 19/02/2026).",
        "  - Plus turnicidad solo con turno Rotativo; plus nocturnidad solo con turno Noche;",
        "    plus penosidad solo con observacion de altura o cristaleria.",
        "  - Zona asignada del empleado = zona de su centro de coste.",
        "  - Todo empleado no-baja tiene asignacion vigente en su centro de coste.",
        "  - Vacaciones sin solapes y <= 30 dias naturales en 2026;",
        "    estado 'Vacaciones' implica estancia aprobada que cubre la fecha de generacion.",
        "  - Próxima revision de maquinaria > fecha de generación; última revisión <= fecha.",
        "  - Auditorias: accion correctiva obligatoria si no es Conforme; cierres <= fecha.",
        "  - Proximas revisiones, garantías y vencimientos siempre relativos a la fecha base.",
        "",
        "Nota: 'Peón' (categoria original) se normalizo a 'Operario' (convenio de limpieza);",
        "'Plus Peligrosidad' se renombro a 'Plus Penosidad'; emails sin tildes ni espacios.",
    ]

    escribir(datos, info)

    print("VALIDACION OK: 0 errores de coherencia")
    for k, v in datos.items():
        print(f"  {k}: {len(v)} filas")
    print(f"Escrito: {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
